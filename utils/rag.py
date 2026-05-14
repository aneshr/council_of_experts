"""
RAG helpers for BYOD ingestion.

This module is responsible for:
- Extracting text from supported document types.
- Chunking text into overlapping segments.
- Embedding chunks with a local Ollama embedding model.
- Persisting / updating a global FAISS index on disk.
"""

import os
import uuid
from datetime import datetime, timezone
from typing import List, Dict, Any, Optional

from langchain_community.embeddings import OllamaEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document

try:
    import PyPDF2  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    PyPDF2 = None

try:
    import docx  # python-docx  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    docx = None


# Directory where the FAISS index will be stored.
_BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_INDEX_DIR = os.path.join(_BASE_DIR, "rag_store", "faiss")
os.makedirs(_INDEX_DIR, exist_ok=True)

_EMBEDDINGS: Optional[OllamaEmbeddings] = None


def _get_embeddings() -> OllamaEmbeddings:
    """
    Lazily initialize and return the Ollama embeddings client.
    """
    global _EMBEDDINGS
    if _EMBEDDINGS is None:
        _EMBEDDINGS = OllamaEmbeddings(
            model="nomic-embed-text",
            base_url="http://localhost:11434",
        )
    return _EMBEDDINGS


def _detect_file_type(filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext.startswith("."):
        ext = ext[1:]
    return ext


def _extract_text_from_txt(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="ignore")


def _extract_text_from_pdf(data: bytes) -> str:
    if PyPDF2 is None:
        raise ValueError(
            "PDF support requires PyPDF2. Install it with `pip install PyPDF2`."
        )
    from io import BytesIO

    reader = PyPDF2.PdfReader(BytesIO(data))
    pages = []
    for page in reader.pages:
        # extract_text may return None
        text = page.extract_text() or ""
        pages.append(text)
    return "\n".join(pages)


def _extract_text_from_docx(data: bytes) -> str:
    if docx is None:
        raise ValueError(
            "DOCX support requires python-docx. Install it with `pip install python-docx`."
        )
    from io import BytesIO

    document = docx.Document(BytesIO(data))
    paragraphs = [p.text for p in document.paragraphs if p.text]
    return "\n".join(paragraphs)


def _extract_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Extract raw text and file_type from the uploaded file.
    """
    file_type = _detect_file_type(filename)

    if file_type in {"txt", "md"}:
        raw_text = _extract_text_from_txt(file_bytes)
    elif file_type == "pdf":
        raw_text = _extract_text_from_pdf(file_bytes)
    elif file_type == "docx":
        raw_text = _extract_text_from_docx(file_bytes)
    else:
        raise ValueError(
            "Unsupported file type. Only .txt, .md, .pdf, .docx are supported for now."
        )

    raw_text = raw_text.strip()
    if not raw_text:
        raise ValueError("The uploaded document contains no extractable text.")

    return {"raw_text": raw_text, "file_type": file_type}


def _chunk_text(
    text: str,
    chunk_size: int = 1500,
    chunk_overlap: int = 300,
) -> List[Dict[str, Any]]:
    """
    Naive character-based text splitter with overlap.

    Returns a list of dicts with:
    - "text": the chunk string
    - "char_start": start index in the original text
    - "char_end": end index in the original text
    """
    chunks: List[Dict[str, Any]] = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = min(start + chunk_size, text_length)
        chunk_text = text[start:end]
        chunks.append(
            {
                "text": chunk_text,
                "char_start": start,
                "char_end": end,
            }
        )
        if end == text_length:
            break
        start = end - chunk_overlap
        if start < 0:
            start = 0

    return chunks


def _load_existing_index(embeddings: OllamaEmbeddings) -> Optional[FAISS]:
    """
    Load an existing FAISS index from disk if it exists.
    """
    try:
        return FAISS.load_local(
            _INDEX_DIR,
            embeddings,
            allow_dangerous_deserialization=True,
        )
    except Exception:
        return None


def ingest_document(
    file_bytes: bytes,
    filename: str,
    doc_name: Optional[str] = None,
    source: str = "upload",
    tags: Optional[List[str]] = None,
    description: Optional[str] = None,
    language: str = "en",
) -> Dict[str, Any]:
    """
    Ingest a single document into the global FAISS index.

    Returns a summary dict that can be sent back from the API layer.
    """
    extracted = _extract_text(file_bytes, filename)
    raw_text: str = extracted["raw_text"]
    file_type: str = extracted["file_type"]

    doc_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    resolved_doc_name = doc_name or filename
    tag_list = tags or []

    chunks = _chunk_text(raw_text)
    documents: List[Document] = []

    for idx, chunk in enumerate(chunks):
        metadata = {
            "doc_id": doc_id,
            "doc_name": resolved_doc_name,
            "filename": filename,
            "source": source,
            "description": description or "",
            "tags": tag_list,
            "file_type": file_type,
            "language": language,
            "ingested_at": now,
            "chunk_index": idx,
            "char_start": chunk["char_start"],
            "char_end": chunk["char_end"],
        }
        documents.append(
            Document(page_content=chunk["text"], metadata=metadata)
        )

    embeddings = _get_embeddings()
    index = _load_existing_index(embeddings)

    if index is None:
        index = FAISS.from_documents(documents, embeddings)
    else:
        index.add_documents(documents)

    index.save_local(_INDEX_DIR)

    return {
        "doc_id": doc_id,
        "doc_name": resolved_doc_name,
        "filename": filename,
        "file_type": file_type,
        "num_chunks": len(chunks),
        "tags": tag_list,
        "ingested_at": now,
    }


def retrieve_context(question: str, k: int = 5) -> List[Document]:
    """
    Retrieve top-k relevant chunks from the global FAISS index.

    Raises ValueError if no index exists yet.
    """
    embeddings = _get_embeddings()
    index = _load_existing_index(embeddings)
    if index is None:
        raise ValueError(
            "No document index found. Ingest at least one document via /ask/byod first."
        )

    # similarity_search returns a list of Documents with .page_content and .metadata
    return index.similarity_search(question, k=k)

